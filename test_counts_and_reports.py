"""
Two numbers that were wrong, and the tests that keep them honest.

  B3  The download button counted every row while the screen and the export
      counted the ACTIVE ones. The header said 706, the button said 707, and the
      file that arrived held 706.
  C5  An import that UPDATED a row reported it as CREATED, so the summary line
      contradicted the per-row detail printed underneath it.
"""
import re
from decimal import Decimal as D
from io import BytesIO

from django.urls import reverse
from openpyxl import Workbook, load_workbook

from accounts.testing import AuthedTestCase
from masters import sheets
from masters.imports import ImportReport
from masters.models import Activity, Material, MaterialGroup, Vendor, VendorGroup


class TheDownloadButtonSaysWhatItWillDownload(AuthedTestCase):
    """
    >>> ANCHOR: MASTER-EXPORT-FOLLOWS-FILTERS <<<
    ⚠⚠ THE BUTTON IS THE SAFEGUARD. The export follows the filters silently, so
       a spreadsheet of 40 rows when you expected the whole master looks like
       data loss. The button stating its own count is the only thing standing
       between a reader and that surprise — which is exactly why its number has
       to be the number, not a different one.
    """

    def setUp(self):
        super().setUp()
        self.rcc = Activity.objects.get(abbreviation="RCC")
        self.group = MaterialGroup.objects.create(code="ZQC", name="ZQ Counts")
        for n in range(3):
            Material.objects.create(
                code=f"ZQC-{n:03d}", name=f"ZQ MATERIAL {n}", group=self.group,
                uom="Nos", estimation_rate=D("10"), gst_percent=D("18"),
                home_activity=self.rcc, is_active=True)
        # ⚠ THE ONE THAT BROKE IT: an inactive row is in the table and is not on
        #   the screen, and was counted by the button anyway.
        self.retired = Material.objects.create(
            code="ZQC-999", name="ZQ RETIRED", group=self.group,
            uom="Nos", estimation_rate=D("10"), gst_percent=D("18"),
            home_activity=self.rcc, is_active=False)

    def test_the_button_count_matches_the_header_count(self):
        page = self.client.get(reverse("material_list")).content.decode()
        active = Material.objects.filter(is_active=True).count()
        self.assertIn(f"{active} materials found", page)
        self.assertIn(f"Download all {active}", page)

    def test_the_button_does_not_count_inactive_rows(self):
        everything = Material.objects.count()
        page = self.client.get(reverse("material_list")).content.decode()
        self.assertNotIn(f"Download all {everything}", page,
                         "the button counted rows the screen does not show — B3")

    def test_the_export_really_returns_the_number_on_the_button(self):
        """The assertion that matters: the count and the file agree."""
        response = self.client.post(reverse("material_export"))
        book = load_workbook(BytesIO(response.content))
        rows = [r for r in book["Materials"].iter_rows(min_row=3, values_only=True)
                if r and r[0]]
        self.assertEqual(len(rows), Material.objects.filter(is_active=True).count())
        self.assertNotIn("ZQC-999", [r[0] for r in rows],
                         "an inactive material reached the export")

    def test_a_filtered_button_says_these_and_counts_the_filter(self):
        page = self.client.get(
            f"{reverse('material_list')}?group=ZQC").content.decode()
        self.assertIn("Download these 3", page)

    def test_vendors_have_the_same_guard(self):
        group = VendorGroup.objects.create(name="ZQ Counts Vendors")
        Vendor.objects.create(code="VEN-ZQ1", name="ZQ Live", phone="9000000001",
                              group=group, is_active=True)
        Vendor.objects.create(code="VEN-ZQ2", name="ZQ Retired", phone="9000000002",
                              group=group, is_active=False)
        page = self.client.get(reverse("vendor_list")).content.decode()
        active = Vendor.objects.filter(is_active=True).count()
        self.assertIn(f"Download all {active}", page)
        self.assertNotIn(f"Download all {Vendor.objects.count()}", page)

    # ---------------------------------------------------------------- C4
    def download_button(self, page):
        """
        The download button's own opening tag.

        ⚠ ASSERT ON THE TAG, NOT ON THE WORD "disabled" ANYWHERE ON THE PAGE.
          The base stylesheet carries `button:disabled{opacity:.4}`, so a bare
          `assertIn("disabled", page)` passes on every screen ever rendered and
          proves nothing. Bug 22's shape: asserting on the page when the claim is
          about one element.
        """
        match = re.search(r"<button[^>]*>\s*(?:{%.*?%}\s*)?[^<]*⤓ Download[^<]*</button>",
                          page, re.S)
        self.assertIsNotNone(match, "no download button on the page at all")
        return match.group(0)

    def test_the_button_is_disabled_when_nothing_matches(self):
        """
        ⚠ C4 — "⤓ Download these 0" was offered and pressable. The empty state
          beside it already says "No materials match that. Clear the filters",
          so the button was contradicting the sentence next to it.
        """
        page = self.client.get(
            f"{reverse('material_list')}?q=zzzznothingmatchesthis").content.decode()
        self.assertIn("Download these 0", page)
        self.assertIn("disabled", self.download_button(page))

    def test_the_button_is_not_disabled_when_something_matches(self):
        """The guard has to be narrow, or every screen ships a dead button."""
        page = self.client.get(
            f"{reverse('material_list')}?group=ZQC").content.decode()
        self.assertIn("Download these 3", page)
        self.assertNotIn("disabled", self.download_button(page))

    def test_vendors_are_disabled_at_zero_too(self):
        page = self.client.get(
            f"{reverse('vendor_list')}?q=zzzznothingmatchesthis").content.decode()
        self.assertIn("Download these 0", page)
        self.assertIn("disabled", self.download_button(page))

    def test_the_blank_template_button_is_never_disabled(self):
        """
        ⚠ THE TWO BUTTONS SIT SIDE BY SIDE AND ONLY ONE DEPENDS ON THE COUNT.
          A blank template has no rows by definition, so disabling it at zero
          matches would take away the one button that still works.
        """
        page = self.client.get(
            f"{reverse('material_list')}?q=zzzznothingmatchesthis").content.decode()
        template = re.search(r"<button[^>]*>[^<]*⤓ Blank template</button>", page, re.S)
        self.assertIsNotNone(template)
        self.assertNotIn("disabled", template.group(0))

    def test_the_export_itself_still_answers_at_zero(self):
        """
        ⚠ THE BUTTON IS A COURTESY, NOT A RULE, and the view is deliberately NOT
          hardened to match. An empty workbook harms nothing — unlike approving
          without a GSTIN — so refusing the POST would be a guard with no danger
          behind it, and one more branch to keep true.
        """
        response = self.client.post(
            f"{reverse('material_export')}?q=zzzznothingmatchesthis")
        self.assertEqual(response.status_code, 200)


class AnUpdateIsReportedAsAnUpdate(AuthedTestCase):
    """
    >>> ANCHOR: IMPORT-REPORT <<<
    ⚠⚠ IT MATTERS MOST ON THE DAY IT MATTERS MOST. Seeding the live server
       imports 705 materials and 174 vendors. "705 created" on a run that
       actually updated 705 existing rows is the difference between "it worked"
       and "I have just overwritten the master".
    """

    def setUp(self):
        super().setUp()
        self.group = VendorGroup.objects.create(name="ZQ Report Group")
        self.vendor = Vendor.objects.create(
            code="VEN-ZQ9", name="ZQ Report Vendor", phone="9000000009",
            group=self.group, payment_terms="")

    def sheet(self, terms, whatsapp=None):
        """
        A one-row vendor workbook — data starts at row 3, as the real one does.

        ⚠ `whatsapp=""` IS THE BLANK-TEMPLATE CASE, and it used to report a
          phantom change. It no longer does — see the class below, which is the
          test that would have caught it.
        """
        book = Workbook()
        page = book.active
        page.title = "Vendors"
        page.append(["Code", "Name", "Contact person", "Phone", "WhatsApp", "GSTIN",
                     "Unregistered", "Group", "Activity 1", "Activity 2",
                     "Document type", "Payment terms", "Address", "Active"])
        page.append(["help"] * 14)
        page.append([self.vendor.code, self.vendor.name, "", self.vendor.phone,
                     f"91{self.vendor.phone}" if whatsapp is None else whatsapp, "",
                     "No", self.group.name, "", "", "PO", terms, "", "Yes"])
        stream = BytesIO()
        book.save(stream)
        stream.seek(0)
        return stream

    def test_changing_a_row_counts_as_updated_not_created(self):
        report = sheets.read_vendors(self.sheet("45"))
        self.assertEqual(len(report.created), 0, "an existing vendor was reported as created")
        self.assertEqual(len(report.updated), 1)
        self.vendor.refresh_from_db()
        self.assertEqual(self.vendor.payment_terms, "45")

    def test_the_summary_line_agrees_with_the_detail_under_it(self):
        report = sheets.read_vendors(self.sheet("45"))
        self.assertIn("0 created", report.summary())
        self.assertIn("1 updated", report.summary())

    def test_a_round_trip_that_changes_nothing_says_so(self):
        sheets.read_vendors(self.sheet("45"))          # first pass applies it
        report = sheets.read_vendors(self.sheet("45"))  # second changes nothing
        self.assertEqual(len(report.created), 0)
        self.assertEqual(len(report.updated), 0)
        self.assertEqual(len(report.skipped), 1)
        self.assertIn("1 unchanged", report.summary())

    def test_touched_gathers_both_kinds_for_the_screen(self):
        report = ImportReport("Test")
        report.create("A", "new")
        report.update("B", "rate")
        self.assertEqual(report.touched, [("A", "new"), ("B", "rate")])

    def test_the_screen_lists_an_update_it_made(self):
        """End to end: the message the person actually reads."""
        response = self.client.post(
            reverse("vendor_import"), {"file": self.sheet("60")}, follow=True)
        body = response.content.decode()
        self.assertIn("1 updated", body)
        self.assertIn("VEN-ZQ9", body)


class AChangeTheModelUndoesIsNotAChange(AnUpdateIsReportedAsAnUpdate):
    """
    >>> ANCHOR: IMPORT-REPORT <<<
    The phantom change, found by accident during the Excel round trip and left
    recorded rather than fixed until now.

    ⚠⚠ THE SHAPE: the importer set `whatsapp_number` to `""` from an empty
       column, `Vendor.save()` immediately refilled it from the phone, and the
       row was reported as "updated: whatsapp_number" having not moved a byte.
       The old check compared the SHEET against the database and answered a
       question nobody asked — what the file proposed, rather than what happened.

    ⚠ SO THE FIX IS NOT ABOUT WHATSAPP. The change list is now read back off the
      object after the save, which is true of any field the model derives for
      itself. A test naming only WhatsApp would pass against a version that
      special-cased WhatsApp and still lied about the next derived field.

    ⚠⚠ BUG 25 — THIS WAS WORKED AROUND IN A TEST FIRST. The helper above used to
       carry a comment explaining that it filled WhatsApp in deliberately to
       avoid the phantom. That comment was the bug report. If a test needs an
       odd workaround, the workaround is the thing to look at.
    """

    def test_a_blank_whatsapp_column_reports_nothing(self):
        sheets.read_vendors(self.sheet("45"))                    # settle the terms
        report = sheets.read_vendors(self.sheet("45", whatsapp=""))
        self.assertEqual(len(report.updated), 0,
                         "the model refilled the field, so nothing actually changed")
        self.assertEqual(len(report.skipped), 1)
        self.assertIn("1 unchanged", report.summary())

    def test_the_number_itself_survives_a_blank_column(self):
        sheets.read_vendors(self.sheet("45", whatsapp=""))
        self.vendor.refresh_from_db()
        self.assertEqual(self.vendor.whatsapp_number, f"91{self.vendor.phone}")

    def test_a_real_change_alongside_a_blank_column_is_still_reported(self):
        """
        ⚠ THE FIX MUST NOT SWALLOW A GENUINE EDIT sharing a row with a derived
          one. Silence would be worse than the phantom it replaced.
        """
        report = sheets.read_vendors(self.sheet("90", whatsapp=""))
        self.assertEqual(len(report.updated), 1)
        detail = dict(report.touched)[self.vendor.code]
        self.assertIn("payment_terms", detail)
        self.assertNotIn("whatsapp_number", detail)

    def test_a_typed_whatsapp_number_is_still_honoured(self):
        """Blank means "derive it". A number means that number, and it is a change."""
        report = sheets.read_vendors(self.sheet("45", whatsapp="919999900001"))
        self.assertEqual(len(report.updated), 1)
        self.vendor.refresh_from_db()
        self.assertEqual(self.vendor.whatsapp_number, "919999900001")

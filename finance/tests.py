"""
The finance screens: they render with data, refuse the wrong role, do not grow
queries with the rows, export what is filtered, and keep the ⓘ panel format.

The money itself is proven in finance/test_ladder.py.
"""
import re
from datetime import date, timedelta
from decimal import Decimal as D
from pathlib import Path

from django.conf import settings
from django.db import connection
from django.test import SimpleTestCase
from django.test.utils import CaptureQueriesContext
from django.urls import reverse

from accounts.models import Role
from masters.models import DocumentType
from projects.bom_models import PurchaseOrder, Receipt
from projects import po_service

from finance import calc, services
from finance.models import RABill, VendorInvoice, VendorPayment
from finance.panels import PANELS
from finance.test_ladder import LadderFixture


class ScreenFixture(LadderFixture):
    """A work order with an approved bill and an open one, and a PO with an invoice."""

    def setUp(self):
        super().setUp()
        self.first = self.bill(self.wo, ["50", "20", "4.5"], ["50", "20", "4.5"])
        services.approve(self.first, self.user)
        self.second = self.bill(self.wo, ["50", "20", "5"])
        self.po = self.make_wo(number="ZQF-PO09", doc_type=DocumentType.PO)
        self.invoice = services.record_vendor_invoice(
            self.po, self.user, "INV-9", date(2026, 9, 1), "1000", "180", "1180")
        services.record_invoice_payment(self.invoice, self.user, "500", paid_on=date.today())
        services.record_advance(self.wo, self.user, "4000", paid_on=date.today())

    def get(self, name, *args, **params):
        response = self.client.get(reverse(name, args=args), params)
        self.assertEqual(response.status_code, 200, name)
        return response.content.decode()


class TheScreensRender(ScreenFixture):

    def test_home(self):
        html = self.get("finance_home")
        self.assertIn("Bills payable now", html)
        self.assertIn(self.vendor.name, html)                # top vendors by balance
        self.assertNotIn("Retention", html)
        self.assertNotIn("{{", html)
        self.assertNotIn("{%", html)

    def test_ra_bill_register_and_filters(self):
        html = self.get("finance_ra_bills")
        self.assertIn(self.first.number, html)
        self.assertIn(self.second.number, html)
        html = self.get("finance_ra_bills", status="approved")
        self.assertIn(self.first.number, html)
        self.assertNotIn(self.second.number, html)
        html = self.get("finance_ra_bills", project=str(self.project.pk), vendor=str(self.vendor.pk))
        self.assertIn(self.first.number, html)
        self.get("finance_ra_bills", project="abc")          # guarded, not a crash

    def test_work_order_screen(self):
        html = self.get("finance_wo", self.wo.pk)
        self.assertIn("49,291", html)                        # certified to date
        self.assertNotIn("Retention", html)                  # switched off: no card, no ledger
        self.assertNotIn("Release retention", html)
        self.assertNotIn("DLP", html)
        self.assertIn(self.second.number, html)
        self.assertIn("is draft", html)                      # the open-bill note
        self.assertIn("4,000", html)                         # the advance paid

    def test_the_bill_screen_open_and_locked(self):
        html = self.get("finance_ra_bill", self.second.pk)
        self.assertIn('name="cert-', html)
        self.assertIn("Certify", html)
        self.assertNotIn("Download PDF", html)
        html = self.get("finance_ra_bill", self.first.pk)
        self.assertNotIn('name="cert-', html)
        self.assertIn("Download PDF", html)
        self.assertIn("48,507", html)                        # net payable
        self.assertNotIn("Less: retention", html)            # zero, so the rung is not printed
        self.assertIn("Record payment", html)
        self.assertIn("locked", html)

    def test_pdf_refused_on_an_open_bill(self):
        response = self.client.get(reverse("finance_ra_bill_pdf", args=[self.second.pk]))
        self.assertEqual(response.status_code, 302)

    def test_pdf_on_an_approved_bill(self):
        try:
            import weasyprint  # noqa: F401
        except ImportError:
            self.skipTest("WeasyPrint is not installed here")
        response = self.client.get(reverse("finance_ra_bill_pdf", args=[self.first.pk]))
        self.assertEqual(response.status_code, 200)
        self.assertEqual(response["Content-Type"], "application/pdf")
        self.assertTrue(response.content.startswith(b"%PDF"))

    def test_po_payments_and_the_forms(self):
        html = self.get("finance_po", self.po.pk)
        self.assertIn("Record a bill", html)
        self.assertIn(self.invoice.vendor_invoice_no, html)
        html = self.get("finance_payments")
        self.assertIn("PV-", html)
        self.assertIn("Excel for Tally", html)
        self.get("finance_ra_bill_pay", self.first.pk)
        self.get("finance_ra_bill_discard", self.second.pk)
        self.get("finance_invoice_new", self.po.pk)
        self.get("finance_po_pay", self.po.pk)
        self.get("finance_po_pay", self.po.pk, kind="advance")

    def test_the_module_is_called_finance_and_accounting_everywhere_a_person_sees_it(self):
        from projects.hub import _TILES
        tile = next(t for t in _TILES if t["key"] == "finance")
        self.assertEqual(tile["title"], "Finance & Accounting")
        self.assertEqual(tile["subtitle"],
                         "Bills against orders, contractor bills, payments and what is owed")
        html = self.get("finance_home")
        self.assertIn("Finance &amp; Accounting", html)
        self.assertNotIn(">Retention<", html)                # no tab
        self.assertNotIn("Vendor invoices", html)
        for tab in ("Overview", "Bills", "RA bills", "Payments", "Vendor ledger", "TDS"):
            self.assertIn(f">{tab}</a>", html, tab)

    def test_finance_po_refuses_a_work_order_and_finance_wo_a_purchase_order(self):
        self.assertEqual(self.client.get(reverse("finance_po", args=[self.wo.pk])).status_code, 404)
        self.assertEqual(self.client.get(reverse("finance_wo", args=[self.po.pk])).status_code, 404)

    def test_po_detail_carries_the_finance_links_and_wo_terms(self):
        html = self.client.get(reverse("po_detail", args=[self.project.pk, self.wo.pk])).content.decode()
        self.assertIn(reverse("finance_wo", args=[self.wo.pk]), html)
        self.assertIn("Mobilisation advance", html)
        self.assertNotIn("Retention", html)                   # switched off: 0% prints nothing
        self.assertNotIn("Defect liability period", html)
        html = self.client.get(reverse("po_detail", args=[self.project.pk, self.po.pk])).content.decode()
        self.assertIn(reverse("finance_po", args=[self.po.pk]), html)
        self.assertNotIn("Mobilisation advance", html)
        draft = self.make_wo(number="ZQF-DRF", approve=False)
        html = self.client.get(reverse("po_detail", args=[self.project.pk, draft.pk])).content.decode()
        self.assertNotIn('name="retention"', html)            # no input any more
        self.assertNotIn('name="dlp_months"', html)
        self.assertIn('name="mobilisation_advance"', html)
        self.assertNotIn("finance/wo/", html)                 # not on a draft
        # An order billed at 10% before the switch-off still says so, read-only.
        old = self.make_wo(number="ZQF-OLD", retention="10")
        html = self.client.get(reverse("po_detail", args=[self.project.pk, old.pk])).content.decode()
        self.assertIn("Retention @ 10", html)

    def test_po_save_writes_the_advance_and_cannot_set_retention(self):
        draft = self.make_wo(number="ZQF-DRF2", approve=False)
        self.client.post(reverse("po_save", args=[self.project.pk, draft.pk]),
                         {"retention": "7.5", "dlp_months": "18", "mobilisation_advance": "5000"})
        draft.refresh_from_db()
        self.assertEqual((draft.retention_pct, draft.dlp_months, draft.mobilisation_advance),
                         (D("0.00"), 12, D("5000.00")))       # the posted retention is ignored

    def test_the_data_migration_clears_ten_percent_only_where_nothing_is_billed(self):
        from importlib import import_module
        from django.apps import apps
        step = import_module("projects.migrations.0012_retention_switched_off")
        billed = self.make_wo(number="ZQF-M1", retention="10")
        bill = self.bill(billed, ["1", "1", "1"], ["1", "1", "1"])
        services.approve(bill, self.user)
        unbilled = self.make_wo(number="ZQF-M2", retention="10")
        typed = self.make_wo(number="ZQF-M3", retention="7.5")
        step.clear_unbilled_defaults(apps, None)
        for order, expected in ((billed, D("10.00")), (unbilled, D("0.00")), (typed, D("7.50"))):
            order.refresh_from_db()
            self.assertEqual(order.retention_pct, expected, order.number)

    def test_a_new_work_order_defaults_to_no_retention(self):
        order = PurchaseOrder.objects.create(number="ZQF-DEF", project=self.project,
                                             vendor=self.vendor, document_type=DocumentType.WO,
                                             created_by=self.user)
        self.assertEqual(order.retention_pct, D("0"))

    def test_po_advance_paid_button_is_refused_on_a_billed_work_order(self):
        self.wo.status = PurchaseOrder.Status.DELIVERED
        self.wo.save()
        self.client.post(reverse("po_advance", args=[self.project.pk, self.wo.pk]), {"to": "paid"})
        self.wo.refresh_from_db()
        self.assertEqual(self.wo.status, PurchaseOrder.Status.DELIVERED)


class TheWritesThroughTheScreens(ScreenFixture):

    def line_ids(self, bill):
        return [line.id for line in bill.lines.order_by("id")]

    def test_save_certify_approve_pay_through_the_forms(self):
        ids = self.line_ids(self.second)
        post = {f"claim-{ids[0]}": "50", f"claim-{ids[1]}": "20", f"claim-{ids[2]}": "5",
                f"cert-{ids[0]}": "50", f"cert-{ids[1]}": "20", f"cert-{ids[2]}": "5",
                "is_final": "1", "contractor_ref": "C-2", "bill_date": "2026-09-02",
                "note": "closing", "action": "certify"}
        self.client.post(reverse("finance_ra_bill_save", args=[self.second.pk]), post)
        self.second.refresh_from_db()
        self.assertEqual(self.second.status, RABill.Status.CERTIFIED)
        self.assertTrue(self.second.is_final)
        self.assertEqual(self.second.contractor_ref, "C-2")

        self.client.post(reverse("finance_ra_bill_approve", args=[self.second.pk]))
        self.second.refresh_from_db()
        self.assertEqual(self.second.status, RABill.Status.APPROVED)
        self.wo.refresh_from_db()
        self.assertEqual(self.wo.status, PurchaseOrder.Status.DELIVERED)
        self.assertEqual(Receipt.objects.filter(source=Receipt.Source.RA_BILL).count(), 6)

        response = self.client.post(reverse("finance_ra_bill_pay", args=[self.second.pk]),
                                    {"amount": "49315.90", "tds_amount": "880.10",
                                     "paid_on": "2026-09-10", "mode": "neft", "reference": "U1"})
        self.assertEqual(response.status_code, 302)
        self.second.refresh_from_db()
        self.assertEqual(self.second.status, RABill.Status.PAID)

    def test_over_certification_through_the_form_leaves_the_bill_a_draft(self):
        ids = self.line_ids(self.second)
        post = {f"cert-{ids[0]}": "50", f"cert-{ids[1]}": "20", f"cert-{ids[2]}": "9",
                "action": "certify"}
        response = self.client.post(reverse("finance_ra_bill_save", args=[self.second.pk]), post,
                                    follow=True)
        self.assertIn("over the order", response.content.decode())
        self.second.refresh_from_db()
        self.assertEqual(self.second.status, RABill.Status.DRAFT)

    def test_new_bill_and_discard_through_the_forms(self):
        services.discard_draft(self.second)
        ids = [line.id for line in self.wo.lines.order_by("id")]
        response = self.client.post(reverse("finance_ra_bill_new", args=[self.wo.pk]),
                                    {f"claim-{ids[0]}": "10", "bill_date": "2026-09-03"})
        bill = RABill.objects.get(purchase_order=self.wo, status=RABill.Status.DRAFT)
        self.assertRedirects(response, reverse("finance_ra_bill", args=[bill.pk]))
        self.assertEqual(bill.lines.count(), 1)
        self.client.post(reverse("finance_ra_bill_discard", args=[bill.pk]))
        self.assertFalse(RABill.objects.filter(pk=bill.pk).exists())

    def test_the_retention_routes_are_gone(self):
        from django.urls import NoReverseMatch
        for name in ("finance_retention", "finance_retention_release", "finance_invoices"):
            with self.assertRaises(NoReverseMatch, msg=name):
                reverse(name, args=[self.wo.pk] if name.endswith("release") else [])
        self.assertEqual(self.client.get("/finance/retention/").status_code, 404)

    def test_invoice_and_payment_forms(self):
        response = self.client.post(reverse("finance_invoice_new", args=[self.po.pk]),
                                    {"vendor_invoice_no": "INV-10", "invoice_date": "2026-09-04",
                                     "taxable": "200", "gst": "36", "total": "236"})
        self.assertRedirects(response, reverse("finance_po", args=[self.po.pk]))
        invoice = VendorInvoice.objects.get(vendor_invoice_no="INV-10")
        response = self.client.post(reverse("finance_invoice_new", args=[self.po.pk]),
                                    {"vendor_invoice_no": "INV-11", "invoice_date": "2026-09-04",
                                     "taxable": "200", "gst": "36", "total": "237"})
        self.assertEqual(response.status_code, 200)          # refused, form redrawn
        self.assertFalse(VendorInvoice.objects.filter(vendor_invoice_no="INV-11").exists())

        self.client.post(reverse("finance_po_pay", args=[self.po.pk]),
                         {"kind": "against_invoice", "invoice": str(invoice.pk), "amount": "236",
                          "paid_on": "2026-09-05", "mode": "upi"})
        self.assertEqual(calc.bulk_invoice_figures([invoice])[invoice.pk]["status"], "settled")
        self.client.post(reverse("finance_po_pay", args=[self.po.pk]),
                         {"kind": "advance", "amount": "300", "paid_on": "2026-09-05", "mode": "cash"})
        self.assertEqual(VendorPayment.objects.filter(purchase_order=self.po,
                                                      kind=VendorPayment.Kind.ADVANCE).count(), 1)


class TheExports(ScreenFixture):

    def sheet(self, response):
        from io import BytesIO
        from openpyxl import load_workbook
        self.assertEqual(response.status_code, 200)
        return load_workbook(BytesIO(response.content))

    def test_ra_bill_excel_carries_the_ladder_and_the_lines(self):
        book = self.sheet(self.client.post(reverse("finance_ra_bills_excel"), {}))
        bills = book["RA bills"]
        rows = list(bills.iter_rows(values_only=True))
        headings = list(rows[0])
        by_number = {row[0]: dict(zip(headings, row)) for row in rows[1:]}
        first = by_number[self.first.number]
        self.assertEqual(D(str(first["Net payable"])), D("48507.40"))
        self.assertNotIn("Retention", headings)               # switched off
        self.assertEqual(D(str(first["Advance recovery"])), D("4900.58"))
        lines = list(book["Lines"].iter_rows(values_only=True))
        self.assertEqual(len(lines), 1 + 6)                  # two bills, three lines each
        # The filter reaches the export.
        book = self.sheet(self.client.post(reverse("finance_ra_bills_excel"), {"status": "draft"}))
        self.assertEqual(len(list(book["RA bills"].iter_rows(values_only=True))), 2)

    def test_payments_excel_has_the_tally_columns(self):
        book = self.sheet(self.client.post(reverse("finance_payments_excel"),
                                           {"from": "2026-01-01", "to": "2030-01-01"}))
        rows = list(book.active.iter_rows(values_only=True))
        self.assertEqual(list(rows[0]), ["Date", "PV number", "Vendor", "GSTIN", "Project",
                                         "Document", "Kind", "Amount", "TDS", "Mode", "Reference"])
        self.assertEqual(len(rows), 1 + VendorPayment.objects.count())


class WhoMayOpenWhat(ScreenFixture):

    def as_role(self, role):
        user = self.make_user(f"zq.{role}", role=role)
        self.client.force_login(user)

    def status(self, name, *args):
        return self.client.get(reverse(name, args=args)).status_code

    def test_site_reaches_the_work_order_and_the_bill_but_not_the_books(self):
        self.as_role(Role.SITE)
        self.assertEqual(self.status("finance_wo", self.wo.pk), 200)
        self.assertEqual(self.status("finance_ra_bill", self.second.pk), 200)
        for name in ("finance_home", "finance_ra_bills", "finance_bills", "finance_bill_pick",
                     "finance_payments", "finance_vendor_ledger", "finance_tds"):
            self.assertEqual(self.status(name), 403, name)
        for name in ("finance_bills_excel", "finance_vendor_ledger_excel", "finance_tds_excel"):
            self.assertEqual(self.client.post(reverse(name), {}).status_code, 403, name)
        self.assertEqual(self.status("finance_po", self.po.pk), 403)
        self.assertEqual(self.status("finance_ra_bill_pdf", self.first.pk), 403)

    def test_purchase_and_compliance_are_refused_everywhere(self):
        for role in (Role.PURCHASE, Role.COMPLIANCE):
            self.as_role(role)
            self.assertEqual(self.status("finance_home"), 403)
            self.assertEqual(self.status("finance_wo", self.wo.pk), 403)
            self.assertEqual(self.status("finance_ra_bill", self.second.pk), 403)

    def test_the_accountant_pays_but_does_not_certify_or_approve(self):
        self.as_role(Role.ACCOUNTANT)
        self.assertEqual(self.status("finance_home"), 200)
        self.assertEqual(self.status("finance_ra_bill_pay", self.first.pk), 200)
        self.assertEqual(self.status("finance_invoice_new", self.po.pk), 200)
        self.assertEqual(self.client.post(reverse("finance_ra_bill_save", args=[self.second.pk]),
                                          {}).status_code, 403)
        self.assertEqual(self.client.post(reverse("finance_ra_bill_approve",
                                                  args=[self.second.pk])).status_code, 403)
        self.assertEqual(self.status("finance_ra_bill_discard", self.second.pk), 403)
        self.assertEqual(self.status("finance_bill_pick"), 200)
        for name in ("finance_bills", "finance_vendor_ledger", "finance_tds"):
            self.assertEqual(self.status(name), 200, name)

    def test_the_project_manager_certifies_and_approves_but_does_not_pay(self):
        self.as_role(Role.PROJECT_MANAGER)
        self.assertEqual(self.status("finance_ra_bill_pay", self.first.pk), 403)
        self.assertEqual(self.status("finance_invoice_new", self.po.pk), 403)
        self.assertEqual(self.status("finance_po_pay", self.po.pk), 403)
        self.assertEqual(self.status("finance_bill_pick"), 403)
        self.assertEqual(self.client.post(reverse("finance_ra_bill_approve",
                                                  args=[self.second.pk])).status_code, 302)
        html = self.client.get(reverse("finance_bills")).content.decode()
        self.assertNotIn("Record a bill", html)               # the button hides with the key

    def test_purchase_and_compliance_are_refused_the_new_registers(self):
        for role in (Role.PURCHASE, Role.COMPLIANCE):
            self.as_role(role)
            for name in ("finance_bills", "finance_vendor_ledger", "finance_tds", "finance_bill_pick"):
                self.assertEqual(self.status(name), 403, f"{role} {name}")

    def test_the_buttons_hide_for_the_role_that_cannot_press_them(self):
        self.as_role(Role.ACCOUNTANT)
        html = self.client.get(reverse("finance_ra_bill", args=[self.second.pk])).content.decode()
        self.assertNotIn("Certify", html)
        self.assertNotIn('name="cert-', html)
        self.as_role(Role.SITE)
        html = self.client.get(reverse("finance_ra_bill", args=[self.first.pk])).content.decode()
        self.assertNotIn("Record payment", html)
        self.assertNotIn("Download PDF", html)


class TheBillsRegister(ScreenFixture):
    """One register, two families — FIN-CALC's bills_register on screen."""

    def test_one_row_per_bill_with_the_same_columns(self):
        rows = calc.bills_register()
        by_number = {r["number"]: r for r in rows}
        self.assertIn(self.invoice.number, by_number)
        self.assertIn(self.first.number, by_number)
        self.assertNotIn(self.second.number, by_number)      # a draft is not a bill
        invoice = by_number[self.invoice.number]
        self.assertEqual((invoice["kind"], invoice["amount"], invoice["paid"], invoice["tds"],
                          invoice["balance"], invoice["status"]),
                         ("po", D("1180.00"), D("500.00"), D("0.00"), D("680.00"), "part_paid"))
        bill = by_number[self.first.number]
        self.assertEqual((bill["kind"], bill["amount"], bill["tds"], bill["paid"], bill["balance"],
                          bill["status"]),
                         ("wo", D("55576.26"), D("862.60"), D("0.00"), D("48507.40"), "open"))
        self.assertEqual(bill["due_date"], date(2026, 9, 1) + timedelta(days=30))
        self.assertTrue(bill["terms_assumed"])               # no number in payment_terms
        totals = calc.bills_total(rows)
        self.assertEqual(totals["balance"], D("49187.40"))
        self.assertEqual(totals["count"], 2)

    def test_a_paid_ra_bill_is_settled(self):
        services.record_ra_payment(self.first, self.user, "48507.40", tds_amount="862.60",
                                   paid_on=date.today())
        row = next(r for r in calc.bills_register() if r["number"] == self.first.number)
        self.assertEqual((row["paid"], row["balance"], row["status"]),
                         (D("48507.40"), D("0.00"), "settled"))

    def test_the_screen_filters_and_totals(self):
        html = self.get("finance_bills")
        self.assertIn(self.invoice.number, html)
        self.assertIn(self.first.number, html)
        self.assertIn("Record a bill", html)
        self.assertIn("49,187", html)                        # tfoot balance
        html = self.get("finance_bills", type="po")
        self.assertIn(self.invoice.number, html)
        self.assertNotIn(self.first.number, html)
        html = self.get("finance_bills", type="wo")
        self.assertNotIn(self.invoice.number, html)
        self.assertIn(self.first.number, html)
        html = self.get("finance_bills", status="part_paid")
        self.assertIn(self.invoice.number, html)
        self.assertNotIn(self.first.number, html)
        html = self.get("finance_bills", project=str(self.project.pk), vendor=str(self.vendor.pk))
        self.assertIn(self.first.number, html)
        html = self.get("finance_bills", **{"from": "2026-09-02"})
        self.assertNotIn(self.first.number, html)            # dated 1 Sep
        self.get("finance_bills", project="abc", type="zz", status="zz")   # guarded

    def test_record_a_bill_goes_through_the_picker_to_the_form(self):
        html = self.get("finance_bill_pick")
        self.assertIn(self.po.number, html)
        self.assertNotIn(self.wo.number, html)               # work orders are billed by RA bills
        self.assertIn(reverse("finance_invoice_new", args=[self.po.pk]), html)
        html = self.get("finance_bill_pick", project=str(self.project.pk))
        self.assertIn(self.po.number, html)
        draft = self.make_wo(number="ZQF-PODR", doc_type=DocumentType.PO, approve=False)
        self.assertNotIn(draft.number, self.get("finance_bill_pick"))

    def test_the_excel_carries_both_families_and_follows_the_filter(self):
        from io import BytesIO
        from openpyxl import load_workbook
        response = self.client.post(reverse("finance_bills_excel"), {})
        self.assertEqual(response.status_code, 200)
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        headings = list(rows[0])
        self.assertEqual(headings[:3], ["Bill no", "Type", "Document"])
        by_number = {row[0]: dict(zip(headings, row)) for row in rows[1:]}
        self.assertEqual(D(str(by_number[self.first.number]["Balance"])), D("48507.40"))
        self.assertEqual(by_number[self.invoice.number]["Type"], "PO")
        response = self.client.post(reverse("finance_bills_excel"), {"type": "po"})
        self.assertEqual(len(list(load_workbook(BytesIO(response.content)).active.rows)), 2)


class TheProjectFilteredTracker(ScreenFixture):

    def test_project_first_and_the_subtotal_row(self):
        html = self.get("finance_ra_bills")
        self.assertNotIn("approved and paid bills ·", html)  # no project chosen → no subtotal
        self.assertIn(">All projects<", html)
        html = self.get("finance_ra_bills", project=str(self.project.pk))
        self.assertIn("certified to date ₹49,291", html)     # approved bills only
        self.assertIn("billed ₹55,576", html)
        self.assertIn("48,507", html)                        # the balance still to pay
        # The tab strip carries the project along, as analytics does.
        self.assertIn(f"{reverse('finance_bills')}?project={self.project.pk}", html)
        self.assertIn(f"{reverse('finance_home')}?project={self.project.pk}", html)
        html = self.get("finance_ra_bills")
        self.assertNotIn(f"{reverse('finance_bills')}?project=", html)

    def test_the_project_select_lists_won_projects(self):
        from projects.models import Project
        lost = Project.objects.create(name="ZQ Lost site", bua_sqft=D("100"),
                                      status=Project.Status.LOST)
        html = self.get("finance_ra_bills")
        self.assertIn(self.project.name, html)
        self.assertNotIn(lost.name, html)


class TheVendorLedger(ScreenFixture):
    """calc.vendor_ledger: bills credit, payments and TDS debit, the balance ties."""

    def test_the_running_balance_ties_to_bills_less_payments(self):
        ledger = calc.vendor_ledger(self.vendor)
        rows = ledger["rows"]
        kinds = [r["kind"] for r in rows]
        self.assertEqual(kinds.count("order"), 2)            # the WO and the PO, both approved
        self.assertEqual(kinds.count("bill"), 2)             # the invoice and the approved RA bill
        self.assertEqual(kinds.count("payment"), 2)          # 500 on the invoice, 4000 advance
        bills = sum((r["bill"] for r in rows), D("0"))
        settled = sum((r["paid"] + r["tds"] for r in rows), D("0"))
        # RA bill row = invoice value − deduction + round off, BEFORE the advance recovery.
        self.assertEqual(bills, D("1180.00") + D("54270.58"))
        self.assertEqual(settled, D("4500.00"))
        self.assertEqual(ledger["closing"], bills - settled)
        self.assertEqual(ledger["closing"], D("50950.58"))
        self.assertTrue(ledger["ties"])
        running = D("0")
        for row in rows:                                     # every row's balance is the running sum
            running += row["bill"] - row["paid"] - row["tds"]
            self.assertEqual(row["balance"], running)
        self.assertTrue(all(r["order_value"] > 0 for r in rows if r["kind"] == "order"))

    def test_the_advance_is_recovered_by_the_bills_and_the_ledger_closes_at_zero(self):
        services.discard_draft(self.second)
        services.record_ra_payment(self.first, self.user, "48507.40", tds_amount="862.60",
                                   paid_on=date(2026, 9, 5))
        final = self.bill(self.wo, ["50", "20", "5"], ["50", "20", "5"], is_final=True)
        services.approve(final, self.user)
        services.record_advance(self.wo, self.user, "6000", paid_on=date(2026, 9, 6))   # 10000 in all
        services.record_ra_payment(final, self.user, "49315.90", tds_amount="880.10",
                                   paid_on=date(2026, 9, 7))
        services.record_invoice_payment(self.invoice, self.user, "680", paid_on=date(2026, 9, 7))
        ledger = calc.vendor_ledger(self.vendor)
        self.assertEqual(ledger["closing"], D("0.00"))

    def test_dates_fold_earlier_rows_into_the_opening_line(self):
        ledger = calc.vendor_ledger(self.vendor, date_from=date(2026, 9, 2))
        self.assertEqual(ledger["opening"]["count"], 2)      # the two bills dated 1 Sep
        self.assertEqual(ledger["opening"]["balance"], D("55450.58"))
        self.assertEqual(ledger["closing"], D("50950.58"))
        ledger = calc.vendor_ledger(self.vendor, date_to=date(2026, 9, 1))
        self.assertEqual(ledger["closing"], D("55450.58"))

    def test_the_screen_and_the_excel(self):
        html = self.get("finance_vendor_ledger")
        self.assertIn("Pick a vendor", html)
        html = self.get("finance_vendor_ledger", vendor=str(self.vendor.pk))
        self.assertIn("50,951", html)                        # closing balance, whole rupees
        self.assertIn("Opening balance", html)
        self.assertIn(self.first.number, html)
        self.assertIn("Excel for Tally", html)
        from io import BytesIO
        from openpyxl import load_workbook
        response = self.client.post(reverse("finance_vendor_ledger_excel"),
                                    {"vendor": str(self.vendor.pk)})
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(rows[0][0], "Date")
        self.assertEqual(rows[1][1], "Opening")
        self.assertEqual(rows[-1][1], "Closing")
        self.assertEqual(D(str(rows[-1][-1])), D("50950.58"))
        self.assertEqual(self.client.post(reverse("finance_vendor_ledger_excel"), {}).status_code, 302)


class TheTDSReport(ScreenFixture):

    def setUp(self):
        super().setUp()
        self.today = date.today()
        services.record_ra_payment(self.first, self.user, "20000", tds_amount="862.60",
                                   paid_on=self.today)
        goods = self.make_wo(number="ZQF-GOODS", doc_type=DocumentType.PO, tds_section="194Q")
        inv = services.record_vendor_invoice(goods, self.user, "G-1", self.today, "2000", "360", "2360")
        services.record_invoice_payment(inv, self.user, "2325", tds_amount="35", paid_on=self.today)

    def test_grouped_by_section_then_vendor_with_the_pro_rated_base(self):
        _key, _label, start, end = calc.quarter_bounds(self.today)
        report = calc.tds_report(start, end)
        self.assertEqual([b["section"] for b in report["sections"]], ["194C", "194Q"])
        contractor = report["sections"][0]
        self.assertEqual(len(contractor["rows"]), 1)         # same vendor, same rate: one row
        row = contractor["rows"][0]
        self.assertEqual(row["count"], 3)                    # 500 on the invoice, 4000 advance, 20000
        self.assertEqual(row["rate"], D("1.75"))
        self.assertEqual(row["tds"], D("862.60"))
        self.assertEqual(row["paid"], D("24500.00"))
        self.assertEqual(row["gross"], D("25362.60"))
        # 1000 × 500 ÷ 1180 = 423.73 · 4000 as itself · 49291.31 × 20862.60 ÷ 49370 = 20829.35
        self.assertEqual(row["base"], D("423.73") + D("4000.00") + D("20829.35"))
        goods = report["sections"][1]["rows"][0]
        self.assertEqual((goods["base"], goods["tds"], goods["paid"]),
                         (D("2000.00"), D("35.00"), D("2325.00")))
        self.assertEqual(report["totals"]["tds"], D("897.60"))
        self.assertEqual(report["totals"]["count"], 4)

    def test_a_quarter_with_nothing_in_it(self):
        report = calc.tds_report(date(2020, 4, 1), date(2020, 6, 30))
        self.assertEqual(report["rows"], [])
        self.assertEqual(report["totals"]["tds"], D("0.00"))

    def test_the_screen_and_the_excel(self):
        html = self.get("finance_tds")
        self.assertIn("194C", html)
        self.assertIn("194Q", html)
        self.assertIn("Excel for 26Q", html)
        self.assertIn("Show quarter", html)
        html = self.get("finance_tds", q="2020-1")             # unknown → the current quarter
        self.assertIn("194C", html)
        from io import BytesIO
        from openpyxl import load_workbook
        response = self.client.post(reverse("finance_tds_excel"), {})
        rows = list(load_workbook(BytesIO(response.content)).active.iter_rows(values_only=True))
        self.assertEqual(list(rows[0])[:3], ["Section", "Vendor", "GSTIN"])
        self.assertEqual(len(rows), 3)                         # heading + two rows


class TheOverview(ScreenFixture):

    def test_kpis_ageing_cash_out_and_top_vendors(self):
        html = self.get("finance_home")
        self.assertIn("Bills payable now", html)
        self.assertIn("49,187", html)                        # 48507.40 + 680
        self.assertIn("Overdue", html)
        self.assertIn("TDS deducted", html)
        self.assertIn(self.vendor.name, html)                # top vendors
        self.assertIn("Ageing of open bills", html)
        self.assertIn("Cash out", html)
        html = self.get("finance_home", project=str(self.project.pk))
        self.assertIn("49,187", html)
        other = self.make_wo(number="ZQF-OTH", doc_type=DocumentType.PO,
                             project=self.other_project())
        services.record_vendor_invoice(other, self.user, "O-1", date.today(), "100", "18", "118")
        self.assertIn("49,305", self.get("finance_home"))    # + 118 across all projects
        self.assertIn("49,187", self.get("finance_home", project=str(self.project.pk)))

    def other_project(self):
        from projects.models import Project
        return Project.objects.create(name="ZQ Other site", bua_sqft=D("100"),
                                      status=Project.Status.WON)

    def test_ageing_buckets_are_by_bill_date(self):
        today = date(2026, 9, 11)
        vendor = self.vendor
        def row(days_old, balance, due_in=None):
            return {"bill_date": today - timedelta(days=days_old), "balance": D(balance),
                    "due_date": today + timedelta(days=due_in if due_in is not None else 30 - days_old),
                    "terms_assumed": False, "vendor": vendor}
        rows = [row(0, "10"), row(30, "20"), row(31, "30"), row(60, "40"), row(61, "50"),
                row(90, "60"), row(91, "70"), row(400, "80"), row(5, "0")]
        buckets = {b["key"]: b for b in calc.ageing(rows, today)}
        self.assertEqual((buckets["d30"]["count"], buckets["d30"]["balance"]), (2, D("30")))
        self.assertEqual((buckets["d60"]["count"], buckets["d60"]["balance"]), (2, D("70")))
        self.assertEqual((buckets["d90"]["count"], buckets["d90"]["balance"]), (2, D("110")))
        self.assertEqual((buckets["older"]["count"], buckets["older"]["balance"]), (2, D("150")))
        late = calc.overdue(rows, today)
        self.assertEqual((late["count"], late["balance"]), (6, D("330")))   # due before today
        windows = {w["key"]: w for w in calc.cash_out(
            [row(0, "1", due_in=0), row(0, "2", due_in=30), row(0, "4", due_in=31),
             row(0, "8", due_in=60), row(0, "16", due_in=90), row(0, "32", due_in=91),
             row(0, "64", due_in=-1)], today)}
        self.assertEqual(windows["w30"]["balance"], D("3"))
        self.assertEqual(windows["w60"]["balance"], D("12"))
        self.assertEqual(windows["w90"]["balance"], D("16"))   # 91 days and overdue left out
        top = calc.top_vendors(rows)
        self.assertEqual((top[0]["count"], top[0]["balance"]), (8, D("360")))


class TheQueryShape(LadderFixture):
    """Query counts must not grow with the rows — FIN-CALC-BULK."""

    def fill(self, count):
        start = PurchaseOrder.objects.count()
        for index in range(start, start + count):
            wo = self.make_wo(number=f"ZQF-Q{index:03d}")
            bill = self.bill(wo, ["100", "40", "10"], ["100", "40", "10"], is_final=True)
            services.approve(bill, self.user)
            services.record_ra_payment(bill, self.user, "1000", paid_on=date.today())
            po = self.make_wo(number=f"ZQF-QP{index:03d}", doc_type=DocumentType.PO)
            invoice = services.record_vendor_invoice(po, self.user, f"I-{index}", date.today(),
                                                     "100", "18", "118")
            services.record_invoice_payment(invoice, self.user, "50", paid_on=date.today())

    def count(self, name, params=None):
        with CaptureQueriesContext(connection) as context:
            response = self.client.get(reverse(name), params or {})
        self.assertEqual(response.status_code, 200, name)
        return len(context.captured_queries)

    def test_registers_do_not_grow_with_the_rows(self):
        self.fill(3)
        screens = ["finance_home", "finance_ra_bills", "finance_bills", "finance_payments",
                   "finance_tds", "finance_bill_pick"]
        params = {"finance_vendor_ledger": {"vendor": str(self.vendor.pk)},
                  "finance_ra_bills": {"project": str(self.project.pk)}}
        small = {name: self.count(name, params.get(name)) for name in screens}
        small["finance_vendor_ledger"] = self.count("finance_vendor_ledger",
                                                    params["finance_vendor_ledger"])
        self.fill(20)
        for name in screens + ["finance_vendor_ledger"]:
            with self.subTest(screen=name):
                self.assertLess(self.count(name, params.get(name)) - small[name], 10)

    def test_the_work_order_and_bill_screens_are_flat(self):
        for index in range(6):
            bill = self.bill(self.wo, ["10", "5", "1"], ["10", "5", "1"])
            services.approve(bill, self.user)
            services.record_ra_payment(bill, self.user, "100", paid_on=date.today())
        with CaptureQueriesContext(connection) as context:
            self.assertEqual(self.client.get(reverse("finance_wo", args=[self.wo.pk])).status_code, 200)
        self.assertLess(len(context.captured_queries), 25)
        with CaptureQueriesContext(connection) as context:
            self.assertEqual(self.client.get(reverse("finance_ra_bill", args=[bill.pk])).status_code, 200)
        self.assertLess(len(context.captured_queries), 25)


class FinancePanelsKeepTheFormat(SimpleTestCase):

    def test_steps_are_between_twenty_and_thirty_words(self):
        for key, panel in PANELS.items():
            for number, entry in enumerate(panel["steps"], 1):
                with self.subTest(panel=key, step=number):
                    self.assertGreaterEqual(len(entry["text"].split()), 20)
                    self.assertLessEqual(len(entry["text"].split()), 30)

    def test_labels_and_step_counts(self):
        for key, panel in PANELS.items():
            with self.subTest(panel=key):
                self.assertTrue(panel["title"])
                self.assertGreaterEqual(len(panel["steps"]), 3)
                for entry in panel["steps"]:
                    self.assertLessEqual(len(entry["label"].split()), 2)
                    self.assertNotIn("<", entry["text"])
                    self.assertNotIn("**", entry["text"])

    def test_perms_name_real_finance_keys(self):
        from accounts.perms import MATRIX
        for key, panel in PANELS.items():
            for entry in panel["steps"]:
                if entry["perm"]:
                    self.assertIn(entry["perm"], MATRIX)
                    self.assertTrue(entry["perm"].startswith("finance."))

    def test_every_key_a_finance_template_uses_exists_and_nothing_more(self):
        folder = Path(settings.BASE_DIR) / "templates" / "finance"
        used = set()
        for path in folder.glob("*.html"):
            used.update(re.findall(r'{% info(?:button|panel) "([^"]+)" %}',
                                   path.read_text(encoding="utf-8")))
        self.assertTrue(used)
        self.assertEqual(used - set(PANELS), set())
        self.assertEqual(set(PANELS) - used, set())
